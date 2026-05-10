import re
import io
import fitz  # PyMuPDF


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def _extract_page_lines(page) -> list[str]:
    """
    Extract PDF text by visual line order instead of relying only on PyMuPDF's
    plain-text mode. Resume PDFs often use columns or positioned text boxes;
    grouping spans by y/x coordinates preserves project titles and bullets more
    reliably for the downstream resume parser.
    """
    page_dict = page.get_text("dict", sort=True)
    rows: list[tuple[float, float, str]] = []

    for block in page_dict.get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            spans = [
                span.get("text", "").strip()
                for span in sorted(line.get("spans", []), key=lambda s: s.get("bbox", [0])[0])
                if span.get("text", "").strip()
            ]
            if not spans:
                continue
            bbox = line.get("bbox", [0, 0, 0, 0])
            rows.append((round(float(bbox[1]), 1), round(float(bbox[0]), 1), " ".join(spans)))

    rows.sort(key=lambda item: (item[0], item[1]))
    return [text for _, _, text in rows]


def extract_from_pdf(file_bytes: bytes) -> str:
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    pages = []
    for page in doc:
        lines = _extract_page_lines(page)
        page_text = "\n".join(lines).strip()
        if not page_text:
            page_text = page.get_text("text")
        pages.append(page_text)
    doc.close()
    return _clean("\n".join(pages))


def extract_from_txt(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "latin-1", "cp1252"):
        try:
            return _clean(file_bytes.decode(encoding))
        except UnicodeDecodeError:
            continue
    raise ValueError("Could not decode text file — unsupported encoding")


def extract_from_docx(file_bytes: bytes) -> str:
    """Extract text from a .docx file using python-docx."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)
        # Also capture text inside tables (common in resume templates)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    if text and text not in paragraphs:
                        paragraphs.append(text)
        return _clean("\n".join(paragraphs))
    except ImportError:
        raise ValueError(
            "DOCX support requires python-docx. Run: pip install python-docx"
        )
    except Exception as exc:
        raise ValueError(f"DOCX extraction failed: {exc}")


def extract_text(file_bytes: bytes, filename: str) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "pdf":
        return extract_from_pdf(file_bytes)
    if ext == "txt":
        return extract_from_txt(file_bytes)
    if ext == "docx":
        return extract_from_docx(file_bytes)
    raise ValueError(
        f"Unsupported file type: .{ext}. Please upload a PDF, DOCX, or TXT file."
    )
